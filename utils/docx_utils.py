from docx.document import Document as Doc
from docx.table import _Cell, Table
from docx.oxml.table import CT_Tbl
from docx.oxml.text.paragraph import CT_P
from docx.text.paragraph import Paragraph
import re

RED = '\033[31m'
RESET = '\033[0m'
def iter_block_items(parent):
    if isinstance(parent, Doc):
        parent_elm = parent.element.body
    elif isinstance(parent, _Cell):
        parent_elm = parent._tc
    else:
        raise ValueError("something's not right")

    for child in parent_elm.iterchildren():
        if isinstance(child, CT_P):
            yield Paragraph(child, parent)

        elif isinstance(child, CT_Tbl):
            yield Table(child, parent)

def get_id(s):
    header = [item.strip() for item in re.split(r'[()（）]', s)]
    # 判断标题中是否包含标识
    test_case_id = None
    for item in header:
        if re.match(r'^[A-Z0-9_]', item):
            test_case_id = item

    if (not test_case_id):
        return s, test_case_id

        # 标题中包含标识
    if (len(header) > 3):
        for i in range(1, 3):
            if i % 2 == 1:
                header[0] += ("（" + header[i])
            else:
                header[0] += ("）" + header[i])
        header[1] = test_case_id

    return header[0], test_case_id

def get_row_content(idx, cols, table):
    row_content = []
    for j in range(cols):
        if j==0 or table.cell(idx, j).text != row_content[-1][0]:
            row_content.append([table.cell(idx, j).text, j])
    return row_content
def get_table_content(table, start = -1, end = -1):
    rows = len(table.rows)
    cols = len(table.columns)
    if start == -1 and end == -1:
        start = 0
        end = rows
    else:
        if end != -1:
            end = min(rows, end)

        if start != -1:
            start = max(start, 0)

    table_content = []
    for i in range(start, end):
        row_content = get_row_content(i, cols, table)
        table_content.append(row_content)

    return table_content

def get_row_cell(idx, cols, table):
    row_cell = []
    for j in range(cols):
        if j==0 or table.cell(idx, j).text != row_cell[-1][0].text:
            row_cell.append([table.cell(idx, j), j])
    return row_cell

def get_table_cell(table, start = -1, end = -1):
    rows = len(table.rows)
    cols = len(table.columns)
    if start == -1 and end == -1:
        start = 0
        end = rows
    else:
        if end != -1:
            end = min(rows, end)

        if start != -1:
            start = max(start, 0)

    table_cell = []
    for i in range(start, end):
        row_cell = get_row_cell(i, cols, table)
        table_cell.append(row_cell)

    return table_cell
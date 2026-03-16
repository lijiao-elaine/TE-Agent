"""
Word文档合并工具
用于生成测试报告汇总文档，支持合并多个Word文档并添加分节符
使用底层ZIP/XML操作确保格式完全保留
"""
import os
from docx import Document
from docx.enum.text import WD_PARAGRAPH_ALIGNMENT
from docx.enum.section import WD_SECTION
from docx.shared import Pt
from docx.oxml.ns import qn
from docx.oxml import parse_xml
from typing import List, Dict, Optional
from copy import deepcopy
import zipfile
import shutil
import tempfile
from xml.etree import ElementTree as ET


class WordDocumentMerger:
    """Word文档合并器，用于生成测试报告汇总文档"""
    
    # Word文档XML命名空间
    NAMESPACES = {
        'w': 'http://schemas.openxmlformats.org/wordprocessingml/2006/main',
        'r': 'http://schemas.openxmlformats.org/officeDocument/2006/relationships',
        'wp': 'http://schemas.openxmlformats.org/drawingml/2006/wordprocessingDrawing',
        'a': 'http://schemas.openxmlformats.org/drawingml/2006/main',
        'pic': 'http://schemas.openxmlformats.org/drawingml/2006/picture',
        'w14': 'http://schemas.microsoft.com/office/word/2010/wordml',
        'mc': 'http://schemas.openxmlformats.org/markup-compatibility/2006',
        'w15': 'http://schemas.microsoft.com/office/word/2012/wordml',
        'wps': 'http://schemas.microsoft.com/office/word/2010/wordprocessingShape',
        'wpc': 'http://schemas.microsoft.com/office/word/2010/wordprocessingCanvas',
        'wpg': 'http://schemas.microsoft.com/office/word/2010/wordprocessingGroup',
    }
    
    def __init__(self, config_manager=None):
        self.config_manager = config_manager
        # 注册命名空间
        for prefix, uri in self.NAMESPACES.items():
            ET.register_namespace(prefix, uri)
    
    def _get_ns(self, prefix):
        """获取命名空间URL"""
        return self.NAMESPACES.get(prefix, '')
    
    def _add_title_paragraph_to_body(self, body, title_text: str):
        """添加标题段落到body元素 - 1级标题，黑体四号"""
        w_ns = self._get_ns('w')
        
        # 创建段落元素
        p = ET.SubElement(body, f'{{{w_ns}}}p')
        
        # 段落属性
        pPr = ET.SubElement(p, f'{{{w_ns}}}pPr')
        # 大纲级别 - 1级
        outlineLvl = ET.SubElement(pPr, f'{{{w_ns}}}outlineLvl')
        outlineLvl.set(f'{{{w_ns}}}val', '0')
        
        # 创建run
        r = ET.SubElement(p, f'{{{w_ns}}}r')
        
        # run属性 - 黑体14号加粗（四号字体 = 14pt）
        rPr = ET.SubElement(r, f'{{{w_ns}}}rPr')
        rFonts = ET.SubElement(rPr, f'{{{w_ns}}}rFonts')
        rFonts.set(f'{{{w_ns}}}eastAsia', '黑体')
        rFonts.set(f'{{{w_ns}}}ascii', '黑体')
        rFonts.set(f'{{{w_ns}}}hAnsi', '黑体')
        rFonts.set(f'{{{w_ns}}}cs', '黑体')
        # 加粗
        b = ET.SubElement(rPr, f'{{{w_ns}}}b')
        bCs = ET.SubElement(rPr, f'{{{w_ns}}}bCs')
        # 字号14pt = 28半磅
        sz = ET.SubElement(rPr, f'{{{w_ns}}}sz')
        sz.set(f'{{{w_ns}}}val', '28')
        szCs = ET.SubElement(rPr, f'{{{w_ns}}}szCs')
        szCs.set(f'{{{w_ns}}}val', '28')
        
        # 文本
        t = ET.SubElement(r, f'{{{w_ns}}}t')
        t.text = title_text
    
    def _add_summary_paragraph_to_body(self, body, software_version: str,
                                       total_cases: int, passed_cases: int,
                                       failed_cases: int, issue_count: int):
        """添加汇总段落到body元素"""
        w_ns = self._get_ns('w')
        
        summary_text = (
            f"针对{software_version}版本的任务目标：完成{software_version}版本测试。"
            f"针对核心测试用例，自动化执行了{total_cases}个测试用例，"
            f"其中{passed_cases}个测试用例测试通过，"
            f"{failed_cases}个测试用例测试未通过，产生{issue_count}个问题单，"
            f"达到阶段性版本冒烟测试的要求，测试执行的具体情况见附录A-测试执行记录，"
            f"测试结果见附录B-测试结果截图，测试中发现的问题报告单见附录C-软件问题报告单。"
        )
        
        # 创建段落元素
        p = ET.SubElement(body, f'{{{w_ns}}}p')
        
        # 段落属性
        pPr = ET.SubElement(p, f'{{{w_ns}}}pPr')
        # 首行缩进480 twips (约2字符)
        ind = ET.SubElement(pPr, f'{{{w_ns}}}ind')
        ind.set(f'{{{w_ns}}}firstLine', '480')
        # 1.5倍行距
        spacing = ET.SubElement(pPr, f'{{{w_ns}}}spacing')
        spacing.set(f'{{{w_ns}}}line', '360')  # 1.5 * 240 = 360
        spacing.set(f'{{{w_ns}}}lineRule', 'auto')
        # 两端对齐
        jc = ET.SubElement(pPr, f'{{{w_ns}}}jc')
        jc.set(f'{{{w_ns}}}val', 'both')
        
        # 创建run
        r = ET.SubElement(p, f'{{{w_ns}}}r')
        
        # run属性 - 宋体小四
        rPr = ET.SubElement(r, f'{{{w_ns}}}rPr')
        rFonts = ET.SubElement(rPr, f'{{{w_ns}}}rFonts')
        rFonts.set(f'{{{w_ns}}}eastAsia', '宋体')
        rFonts.set(f'{{{w_ns}}}ascii', '宋体')
        rFonts.set(f'{{{w_ns}}}hAnsi', '宋体')
        sz = ET.SubElement(rPr, f'{{{w_ns}}}sz')
        sz.set(f'{{{w_ns}}}val', '24')  # 12pt * 2 = 24 half-points
        szCs = ET.SubElement(rPr, f'{{{w_ns}}}szCs')
        szCs.set(f'{{{w_ns}}}val', '24')
        
        # 文本
        t = ET.SubElement(r, f'{{{w_ns}}}t')
        t.text = summary_text
    
    def _add_title_paragraph(self, doc: Document, title_text: str, font_size: int = 14, bold: bool = True):
        """添加标题段落（使用python-docx方式）"""
        para = doc.add_paragraph()
        run = para.add_run(title_text)
        run.font.size = Pt(font_size)
        run.font.bold = bold
        run.font.name = '黑体'
        try:
            if run._element.rPr is None:
                run._element.get_or_add_rPr()
            if run._element.rPr.rFonts is None:
                run._element.rPr._append(parse_xml(f'<w:rFonts xmlns:w="http://schemas.openxmlformats.org/wordprocessingml/2006/main" w:eastAsia="黑体"/>'))
            else:
                run._element.rPr.rFonts.set(qn('w:eastAsia'), '黑体')
        except Exception:
            pass
        para.alignment = WD_PARAGRAPH_ALIGNMENT.LEFT
        return para
    
    def _add_section_break_to_body(self, body):
        """添加分节符到body"""
        w_ns = self._get_ns('w')
        
        # 创建包含分节符的段落
        p = ET.SubElement(body, f'{{{w_ns}}}p')
        pPr = ET.SubElement(p, f'{{{w_ns}}}pPr')
        sectPr = ET.SubElement(pPr, f'{{{w_ns}}}sectPr')
        
        # 分节符类型 - 下一页
        type_elem = ET.SubElement(sectPr, f'{{{w_ns}}}type')
        type_elem.set(f'{{{w_ns}}}val', 'nextPage')
        
        # 页面大小 A4
        pgSz = ET.SubElement(sectPr, f'{{{w_ns}}}pgSz')
        pgSz.set(f'{{{w_ns}}}w', '11906')
        pgSz.set(f'{{{w_ns}}}h', '16838')
        
        # 页边距
        pgMar = ET.SubElement(sectPr, f'{{{w_ns}}}pgMar')
        pgMar.set(f'{{{w_ns}}}top', '1440')
        pgMar.set(f'{{{w_ns}}}right', '1800')
        pgMar.set(f'{{{w_ns}}}bottom', '1440')
        pgMar.set(f'{{{w_ns}}}left', '1800')
        pgMar.set(f'{{{w_ns}}}header', '720')
        pgMar.set(f'{{{w_ns}}}footer', '720')
        pgMar.set(f'{{{w_ns}}}gutter', '0')
    
    def _copy_xml_element_preserving_ns(self, src_elem):
        """复制XML元素，保留所有命名空间"""
        # 创建新的元素，使用相同的标签和命名空间
        tag = src_elem.tag
        
        # 复制属性
        attrib = dict(src_elem.attrib)
        
        # 创建新元素
        new_elem = ET.Element(tag, attrib)
        new_elem.text = src_elem.text
        new_elem.tail = src_elem.tail
        
        # 递归复制子元素
        for child in src_elem:
            new_child = self._copy_xml_element_preserving_ns(child)
            new_elem.append(new_child)
        
        return new_elem
    
    def _merge_document_content(self, base_body, base_temp_dir, append_doc_path, add_section=False, title=None):
        """合并单个文档内容到基础body，并复制媒体文件"""
        w_ns = self._get_ns('w')
        r_ns = self._get_ns('r')
        
        # 解压追加文档
        with tempfile.TemporaryDirectory() as temp_dir:
            with zipfile.ZipFile(append_doc_path, 'r') as zip_ref:
                zip_ref.extractall(temp_dir)
            
            # 读取document.xml
            doc_xml_path = os.path.join(temp_dir, 'word', 'document.xml')
            if not os.path.exists(doc_xml_path):
                return
            
            # 解析XML
            tree = ET.parse(doc_xml_path)
            root = tree.getroot()
            
            # 获取body
            body = root.find(f'.//{{{w_ns}}}body')
            if body is None:
                return
            
            # 复制媒体文件，获取文件名映射（处理重名文件）
            media_mapping = {}
            self._copy_media_files(temp_dir, base_temp_dir, media_mapping)
            
            # 复制关系文件（用于图片引用），获取ID映射，同时更新媒体文件路径
            id_mapping = self._copy_relationships(temp_dir, base_temp_dir, media_mapping)
            
            # 如果需要，添加分节符
            if add_section:
                self._add_section_break_to_body(base_body)
            
            # 如果需要，添加标题
            if title:
                self._add_title_paragraph_to_body(base_body, title)
            
            # 复制body中的所有子元素（除了sectPr）
            for child in body:
                if child.tag == f'{{{w_ns}}}sectPr':
                    continue
                
                # 深拷贝元素
                new_child = deepcopy(child)
                
                # 如果存在ID映射，更新图片引用
                if id_mapping:
                    self._update_image_references(new_child, id_mapping)
                
                base_body.append(new_child)
    
    def _copy_styles_xml(self, source_doc_path, target_temp_dir):
        """复制样式文件"""
        source_styles_path = os.path.join(source_doc_path, 'word', 'styles.xml')
        if os.path.exists(source_styles_path):
            target_styles_path = os.path.join(target_temp_dir, 'word', 'styles.xml')
            shutil.copy2(source_styles_path, target_styles_path)
    
    def _copy_numbering_xml(self, source_doc_path, target_temp_dir):
        """复制编号定义文件"""
        source_numbering_path = os.path.join(source_doc_path, 'word', 'numbering.xml')
        if os.path.exists(source_numbering_path):
            target_numbering_path = os.path.join(target_temp_dir, 'word', 'numbering.xml')
            shutil.copy2(source_numbering_path, target_numbering_path)
    
    def _copy_fonts_table_xml(self, source_doc_path, target_temp_dir):
        """复制字体表文件"""
        source_fonts_path = os.path.join(source_doc_path, 'word', 'fontTable.xml')
        if os.path.exists(source_fonts_path):
            target_fonts_path = os.path.join(target_temp_dir, 'word', 'fontTable.xml')
            shutil.copy2(source_fonts_path, target_fonts_path)
    
    def _copy_settings_xml(self, source_doc_path, target_temp_dir):
        """复制设置文件"""
        source_settings_path = os.path.join(source_doc_path, 'word', 'settings.xml')
        if os.path.exists(source_settings_path):
            target_settings_path = os.path.join(target_temp_dir, 'word', 'settings.xml')
            shutil.copy2(source_settings_path, target_settings_path)
    
    def _copy_theme_files(self, source_doc_path, target_temp_dir):
        """复制主题文件"""
        source_theme_dir = os.path.join(source_doc_path, 'word', 'theme')
        if os.path.exists(source_theme_dir):
            target_theme_dir = os.path.join(target_temp_dir, 'word', 'theme')
            if not os.path.exists(target_theme_dir):
                os.makedirs(target_theme_dir)
            for filename in os.listdir(source_theme_dir):
                source_file = os.path.join(source_theme_dir, filename)
                target_file = os.path.join(target_theme_dir, filename)
                shutil.copy2(source_file, target_file)
    
    def _copy_media_files(self, source_temp_dir, target_temp_dir, media_mapping=None):
        """复制媒体文件（图片等）从源临时目录到目标临时目录
        :param media_mapping: 用于返回文件名映射关系的字典 {原文件名: 新文件名}
        """
        source_media_dir = os.path.join(source_temp_dir, 'word', 'media')
        if os.path.exists(source_media_dir):
            target_media_dir = os.path.join(target_temp_dir, 'word', 'media')
            if not os.path.exists(target_media_dir):
                os.makedirs(target_media_dir)
            
            existing_files = set(os.listdir(target_media_dir))
            
            for filename in os.listdir(source_media_dir):
                source_file = os.path.join(source_media_dir, filename)
                target_file = os.path.join(target_media_dir, filename)
                new_filename = filename
                
                # 如果文件名已存在，生成新文件名
                if filename in existing_files:
                    # 生成新的唯一文件名
                    name, ext = os.path.splitext(filename)
                    counter = 1
                    while new_filename in existing_files:
                        new_filename = f"{name}_{counter}{ext}"
                        counter += 1
                    target_file = os.path.join(target_media_dir, new_filename)
                    
                    # 记录映射关系
                    if media_mapping is not None:
                        media_mapping[filename] = new_filename
                
                shutil.copy2(source_file, target_file)
                existing_files.add(new_filename)
    
    def _copy_relationships(self, source_temp_dir, target_temp_dir, media_mapping=None):
        """复制关系文件（document.xml.rels），返回ID映射字典
        :param media_mapping: 媒体文件名映射 {原文件名: 新文件名}
        """
        source_rels_dir = os.path.join(source_temp_dir, 'word', '_rels')
        id_mapping = {}  # 记录旧ID到新ID的映射
        if os.path.exists(source_rels_dir):
            target_rels_dir = os.path.join(target_temp_dir, 'word', '_rels')
            if not os.path.exists(target_rels_dir):
                os.makedirs(target_rels_dir)
            
            source_rels_file = os.path.join(source_rels_dir, 'document.xml.rels')
            target_rels_file = os.path.join(target_rels_dir, 'document.xml.rels')
            
            if os.path.exists(source_rels_file):
                if os.path.exists(target_rels_file):
                    # 合并关系文件，获取ID映射，同时更新媒体文件路径
                    id_mapping = self._merge_relationships_files(source_rels_file, target_rels_file, media_mapping)
                else:
                    # 复制并更新媒体文件路径
                    if media_mapping:
                        self._update_media_paths_in_rels(source_rels_file, target_rels_file, media_mapping)
                    else:
                        shutil.copy2(source_rels_file, target_rels_file)
        return id_mapping
    
    def _update_media_paths_in_rels(self, source_rels_file, target_rels_file, media_mapping):
        """复制关系文件并更新媒体文件路径"""
        try:
            tree = ET.parse(source_rels_file)
            root = tree.getroot()
            
            for rel in root:
                target = rel.get('Target', '')
                # 检查是否是媒体文件引用
                if target.startswith('media/'):
                    filename = target.replace('media/', '')
                    if filename in media_mapping:
                        # 更新为新的文件名
                        new_target = f"media/{media_mapping[filename]}"
                        rel.set('Target', new_target)
            
            tree.write(target_rels_file, encoding='UTF-8', xml_declaration=True)
        except Exception as e:
            print(f"更新关系文件时出错: {e}")
            shutil.copy2(source_rels_file, target_rels_file)
    
    def _merge_relationships_files(self, source_rels_file, target_rels_file, media_mapping=None):
        """合并两个关系文件，避免ID冲突，返回旧ID到新ID的映射
        :param media_mapping: 媒体文件名映射 {原文件名: 新文件名}
        """
        id_mapping = {}  # 旧ID -> 新ID
        try:
            # 解析源关系文件
            source_tree = ET.parse(source_rels_file)
            source_root = source_tree.getroot()
            
            # 解析目标关系文件
            target_tree = ET.parse(target_rels_file)
            target_root = target_tree.getroot()
            
            # 获取目标文件中已有的rId
            existing_ids = set()
            for rel in target_root:
                rid = rel.get('Id')
                if rid:
                    existing_ids.add(rid)
            
            # 复制源文件中的关系到目标文件
            for rel in source_root:
                # 生成新的唯一ID
                original_id = rel.get('Id', '')
                new_id = original_id
                counter = 1
                while new_id in existing_ids:
                    new_id = f"{original_id}_{counter}"
                    counter += 1
                
                # 记录ID映射
                if new_id != original_id:
                    id_mapping[original_id] = new_id
                
                # 更新媒体文件路径（如果有映射）
                if media_mapping:
                    target_path = rel.get('Target', '')
                    if target_path.startswith('media/'):
                        filename = target_path.replace('media/', '')
                        if filename in media_mapping:
                            rel.set('Target', f"media/{media_mapping[filename]}")
                
                # 更新ID并添加到目标文件
                rel.set('Id', new_id)
                target_root.append(rel)
                existing_ids.add(new_id)
            
            # 保存合并后的文件
            target_tree.write(target_rels_file, encoding='UTF-8', xml_declaration=True)
        except Exception as e:
            print(f"合并关系文件时出错: {e}")
            # 如果合并失败，直接复制
            shutil.copy2(source_rels_file, target_rels_file)
        
        return id_mapping
    
    def _update_image_references(self, element, id_mapping):
        """递归更新XML元素中的图片引用ID"""
        r_ns = self._get_ns('r')
        
        # 定义可能的图片引用属性
        ref_attrs = [
            f'{{{r_ns}}}embed',  # a:blip 的 r:embed
            f'{{{r_ns}}}link',   # a:blip 的 r:link
            f'{{{r_ns}}}id',     # v:imagedata 的 r:id
        ]
        
        # 检查当前元素是否有图片引用属性
        for attr in ref_attrs:
            old_id = element.get(attr)
            if old_id and old_id in id_mapping:
                element.set(attr, id_mapping[old_id])
        
        # 递归处理子元素
        for child in element:
            self._update_image_references(child, id_mapping)
    
    def create_summary_report(self, output_path: str,
                             software_version: str,
                             total_cases: int,
                             passed_cases: int,
                             failed_cases: int,
                             issue_count: int,
                             documents_to_append: List[Dict[str, str]]) -> str:
        """
        创建汇总报告文档
        """
        w_ns = self._get_ns('w')
        
        with tempfile.TemporaryDirectory() as temp_dir:
            # 使用第一个要追加的文档作为基础（如果存在），否则创建新的
            base_doc_path = None
            for doc_info in documents_to_append:
                if os.path.exists(doc_info.get('path', '')):
                    base_doc_path = doc_info.get('path')
                    break
            
            if base_doc_path is None:
                # 创建一个新的空文档作为基础
                empty_doc = Document()
                base_doc_path = os.path.join(temp_dir, 'base.docx')
                empty_doc.save(base_doc_path)
            
            # 解压基础文档
            work_dir = os.path.join(temp_dir, 'work')
            with zipfile.ZipFile(base_doc_path, 'r') as zip_ref:
                zip_ref.extractall(work_dir)
            
            # 读取document.xml
            doc_xml_path = os.path.join(work_dir, 'word', 'document.xml')
            tree = ET.parse(doc_xml_path)
            root = tree.getroot()
            
            # 获取body
            body = root.find(f'.//{{{w_ns}}}body')
            if body is None:
                raise ValueError("无法找到文档body")
            
            # 移除现有的所有内容，保留body结构
            # 先找到sectPr
            existing_sectPr = body.find(f'{{{w_ns}}}sectPr')
            
            # 清空body（保留sectPr在最后）
            for child in list(body):
                if child != existing_sectPr:
                    body.remove(child)
            
            # 添加汇总标题
            self._add_title_paragraph_to_body(body, "1 详细测试结果")
            
            # 添加汇总段落
            self._add_summary_paragraph_to_body(body, software_version, total_cases,
                                                passed_cases, failed_cases, issue_count)
            
            # 合并每个文档
            for i, doc_info in enumerate(documents_to_append):
                doc_path = doc_info.get('path', '')
                title = doc_info.get('title', None)
                
                if not os.path.exists(doc_path):
                    continue
                
                # 第一个文档：如果有标题则加分节符，否则不加
                # 后续文档：都加分节符
                add_section = (i > 0) or (title is not None)
                
                self._merge_document_content(body, work_dir, doc_path, add_section, title)
            
            # 添加最终的sectPr
            if existing_sectPr is not None:
                body.append(existing_sectPr)
            else:
                # 创建默认sectPr
                sectPr = ET.SubElement(body, f'{{{w_ns}}}sectPr')
                pgSz = ET.SubElement(sectPr, f'{{{w_ns}}}pgSz')
                pgSz.set(f'{{{w_ns}}}w', '11906')
                pgSz.set(f'{{{w_ns}}}h', '16838')
                pgMar = ET.SubElement(sectPr, f'{{{w_ns}}}pgMar')
                pgMar.set(f'{{{w_ns}}}top', '1440')
                pgMar.set(f'{{{w_ns}}}right', '1800')
                pgMar.set(f'{{{w_ns}}}bottom', '1440')
                pgMar.set(f'{{{w_ns}}}left', '1800')
                pgMar.set(f'{{{w_ns}}}header', '720')
                pgMar.set(f'{{{w_ns}}}footer', '720')
                pgMar.set(f'{{{w_ns}}}gutter', '0')
            
            # 保存document.xml
            tree.write(doc_xml_path, encoding='UTF-8', xml_declaration=True)
            
            # 确保输出目录存在
            output_dir = os.path.dirname(output_path)
            if output_dir and not os.path.exists(output_dir):
                os.makedirs(output_dir)
            
            # 重新打包为docx
            with zipfile.ZipFile(output_path, 'w', zipfile.ZIP_DEFLATED) as zipf:
                for root_dir, dirs, files in os.walk(work_dir):
                    for file in files:
                        file_path = os.path.join(root_dir, file)
                        arcname = os.path.relpath(file_path, work_dir)
                        zipf.write(file_path, arcname)
            
            print(f"汇总报告已生成: {output_path}")
            return output_path


def generate_all_in_one_report(config_manager, test_stats: Dict[str, int]) -> str:
    """
    生成汇总报告（All-in-One）的便捷函数
    """
    merger = WordDocumentMerger(config_manager)
    
    # 获取配置
    software_version = config_manager.get_software_version()
    output_path = config_manager.get_test_report_allinone_file()
    
    # 准备要追加的文档
    documents_to_append = []
    
    # 1. 测试执行记录
    result_word_file = config_manager.get_result_word_file()
    if os.path.exists(result_word_file):
        documents_to_append.append({
            'path': result_word_file,
            'title': '附录A-测试执行记录'
        })
    
    # 2. 测试结果截图
    screenshots_file = config_manager.get_screenshots_output_file()
    if os.path.exists(screenshots_file):
        documents_to_append.append({
            'path': screenshots_file,
            'title': '附录B-测试结果截图'
        })
    
    # 3. 问题报告单
    issue_report_file = config_manager.get_issue_report_output()
    if os.path.exists(issue_report_file):
        documents_to_append.append({
            'path': issue_report_file,
            'title': '附录C-软件问题报告单'  # 统一添加一级标题
        })
    
    if not documents_to_append:
        print("没有可追加的文档，跳过生成汇总报告")
        return None
    
    return merger.create_summary_report(
        output_path=output_path,
        software_version=software_version,
        total_cases=test_stats.get('total_cases', 0),
        passed_cases=test_stats.get('passed_cases', 0),
        failed_cases=test_stats.get('failed_cases', 0),
        issue_count=test_stats.get('issue_count', 0),
        documents_to_append=documents_to_append
    )
